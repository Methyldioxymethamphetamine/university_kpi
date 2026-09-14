"""Load the REAL, corrected 229-row KPI dictionary
(`registry/client-docs/kpi_dictionary_v2.xlsx`) for discovery/candidates.py.

Supersedes derived/kpi_dictionary_partial.yaml's 24-code stopgap, which was
built only because this file did not exist in the repo yet -- see that
file's own superseded-by note.

**Never load `DOC-20260901-WA0019.xlsx` (the original) from here.** It has
all 229 rows marked `Benchmark_Direction = "Higher"`, including the three
rank KPIs (X03/X04/X05) where that is inverted (rank 1 beats rank 200) --
using it would silently reintroduce a known, already-corrected defect into
every direction-sensitive KPI comparison downstream of this module.

The label text used for token-overlap scoring is `Scraping_Keywords`
(already `Sub_Parameter, Variable_Name, Definition` concatenated --
notes/methodology.md Sec3.5 verified this is 229/229 character-exact and
contributes zero tokens beyond those three columns), not a hand-picked
single field -- this gives discovery the richest text this dictionary
actually offers, not a paraphrase.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
V2_DICTIONARY_PATH = REPO_ROOT / "registry" / "client-docs" / "kpi_dictionary_v2.xlsx"
EXPECTED_ROW_COUNT = 229
EXPECTED_DIRECTION_DISTRIBUTION = {
    "Higher": 205,
    "Lower": 17,
    "Lower (absolute variance)": 1,
    "Context": 3,
    "N/A": 3,
}


@dataclass(frozen=True)
class KpiEntry:
    kpi_code: str
    label: str  # Scraping_Keywords -- see module docstring
    variable_name: str
    definition: str
    unit: str | None
    benchmark_direction: str | None
    contestable: bool
    v2_note: str | None


class DictionaryConsistencyError(RuntimeError):
    """Raised when kpi_dictionary_v2.xlsx does not match its own claimed
    shape -- callers must not silently proceed on an unverified dictionary
    (this session's own rule, restated from the human instruction that
    created this loader)."""


def load_kpi_dictionary_v2(path: Path = V2_DICTIONARY_PATH) -> dict[str, KpiEntry]:
    if not path.exists():
        raise FileNotFoundError(
            f"{path} not found -- kpi_dictionary_v2.xlsx must be placed in "
            f"registry/client-docs/ before discovery can run against the real dictionary"
        )
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["KPI_Dictionary"]
    header = [c.value for c in ws[1]]
    raw_rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    raw_rows = [r for r in raw_rows if r.get("KPI_Code")]

    if len(raw_rows) != EXPECTED_ROW_COUNT:
        raise DictionaryConsistencyError(
            f"expected {EXPECTED_ROW_COUNT} KPI_Code rows, found {len(raw_rows)}"
        )

    from collections import Counter
    actual_distribution = dict(Counter(r.get("Benchmark_Direction") for r in raw_rows))
    if actual_distribution != EXPECTED_DIRECTION_DISTRIBUTION:
        raise DictionaryConsistencyError(
            f"Benchmark_Direction distribution mismatch.\n"
            f"  expected: {EXPECTED_DIRECTION_DISTRIBUTION}\n"
            f"  actual:   {actual_distribution}"
        )

    entries: dict[str, KpiEntry] = {}
    for r in raw_rows:
        v2_note = r.get("v2_Note")
        entries[r["KPI_Code"]] = KpiEntry(
            kpi_code=r["KPI_Code"],
            label=r.get("Scraping_Keywords") or f"{r.get('Sub_Parameter', '')}, {r.get('Variable_Name', '')}, {r.get('Definition', '')}",
            variable_name=r.get("Variable_Name") or "",
            definition=r.get("Definition") or "",
            unit=r.get("Unit"),
            benchmark_direction=r.get("Benchmark_Direction"),
            contestable=bool(v2_note and "CONTESTABLE" in v2_note),
            v2_note=v2_note,
        )
    return entries


def verify_consistency(path: Path = V2_DICTIONARY_PATH) -> dict:
    """Standalone check callable before anything else touches the file --
    returns a small report dict rather than just True/False, so a caller
    can print exactly what was checked, not just that it passed."""
    entries = load_kpi_dictionary_v2(path)  # raises DictionaryConsistencyError on mismatch
    contestable = sorted(code for code, e in entries.items() if e.contestable)
    return {
        "row_count": len(entries),
        "expected_row_count": EXPECTED_ROW_COUNT,
        "contestable_codes": contestable,
    }
